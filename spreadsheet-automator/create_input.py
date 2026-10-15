# spreadsheet-automator/create_input.py
# Helper script: creates the sample input file data/input_sales.xlsx
import os
import openpyxl
from openpyxl.styles import Font

rows = [
    ["Product Name", "Quantity", "Unit Price"],
    ["Laptop", 2, 1200.50],
    ["Mouse", 5, 25.00],
    ["Keyboard", 3, 75.99],
    ["Monitor", 1, 300.00],
    [None, None, None],  # Empty row to show the script skips it
    ["Printer", 1, 150.00],
]

base = os.path.dirname(os.path.abspath(__file__))
path = os.path.join(base, "data", "input_sales.xlsx")
os.makedirs(os.path.dirname(path), exist_ok=True)

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Sheet1"
for r in rows:
    ws.append(r)
for cell in ws[1]:
    cell.font = Font(bold=True)
for row in ws.iter_rows(min_row=2, min_col=3, max_col=3):
    for cell in row:
        cell.number_format = '0.00'
ws.column_dimensions['A'].width = 16
ws.column_dimensions['B'].width = 10
ws.column_dimensions['C'].width = 12

wb.save(path)
print(f"Sample input created: {path}")
